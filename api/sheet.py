"""Reads a screening-camp list (CSV or Excel) into a header row and data rows, all as text.

Matching the columns to the model's inputs happens in the browser, which also scores each row; nothing is stored.
"""

from __future__ import annotations

import base64
import csv
import io

MAX_BYTES = 2 * 1024 * 1024
MAX_ROWS = 200


class SheetError(Exception):
    """A problem the user can act on (shown as is)."""


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_sheet(name: str, data: str) -> dict:
    try:
        raw = base64.b64decode(data, validate=True)
    except ValueError as error:
        raise SheetError(f"{name} could not be read.") from error
    if len(raw) > MAX_BYTES:
        raise SheetError(f"{name} is larger than {MAX_BYTES // 1024 // 1024} MB.")
    lower = name.lower()
    if lower.endswith(".csv"):
        text = raw.decode("utf-8-sig", errors="replace")
        rows = [row for row in csv.reader(io.StringIO(text)) if any(cell.strip() for cell in row)]
    elif lower.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        try:
            sheet = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).worksheets[0]
        except Exception as error:  # not a valid workbook
            raise SheetError(f"{name} is not a readable Excel file.") from error
        rows = [[_text(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
        rows = [row for row in rows if any(row)]
    else:
        raise SheetError("Use a CSV file or an Excel workbook (.xlsx).")
    if len(rows) < 2:
        raise SheetError("The file needs a header row and at least one patient.")
    if len(rows) - 1 > MAX_ROWS:
        raise SheetError(f"Score at most {MAX_ROWS} patients at a time.")
    header = [_text(cell) for cell in rows[0]]
    return {"columns": header, "rows": [[_text(cell) for cell in row] + [""] * (len(header) - len(row))
                                        for row in rows[1:]]}
